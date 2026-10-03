"""Self-test: human() typesets correctly, a clean build passes check(), and every check rule fires.
Run: python tests/selftest.py        (set DOCB_WORD=1 to also run Word finalize + PDF checks on Windows)"""
import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from docx import Document  # noqa: E402
from docbuilder import DesignDoc, check, finalize, human  # noqa: E402

tmp = Path(tempfile.mkdtemp())
ok = True


def expect(name, cond):
    global ok
    ok &= bool(cond)
    print(("PASS " if cond else "FAIL ") + name)


# 1. human() typesetting
expect("middle dot -> comma", human("P1v3 · Linux") == "P1v3, Linux")
expect("short IP -> brackets", human("pep-kv · .66.12") == "pep-kv (.66.12)")
expect("plus -> and", human("DRS 2.1 + Bot Manager") == "DRS 2.1 and Bot Manager")
expect("curly apostrophe", human("Microsoft's") == "Microsoft’s")
expect("curly single quotes", human("the 'testing' rule") == "the ‘testing’ rule")

# 2. a clean build from a blank template passes every check
tpl = tmp / "tpl.docx"
Document().save(tpl)
png = tmp / "fig.png"
try:
    from PIL import Image
    Image.new("RGB", (1600, 800), "white").save(png)
except ImportError:
    png = None
d = DesignDoc(tpl, tmp / "clean.docx")
d.clear_body()
d.h1("Purpose")
d.para("This document describes the proposed design. It records the configuration as of 24-Sep-2026.")
d.h1("Design")
d.h2("Pattern")
d.para("Traffic enters through the WAF. Section 2.1 explains the pattern.")
d.bullets(["The web app keeps public access disabled.", "Rules reference application security groups."])
d.table(["Step", "Path"], [["1", "Internet → WAF → private endpoint"]], [10, 90], title="Table 1: Flow")
if png:
    d.figure(png, "Figure 1: Architecture", heading="Architecture")
d.h1("Plan")
d.para("The cloud team applies the changes in a change window.")
clean = d.save()
found = check(clean)
expect("clean document passes check", found == [])
if found:
    print("   ", *found, sep="\n    ")

# 3. each rule fires on a seeded defect
bad = Document()
bad.add_heading("Findings", 1)
bad.add_paragraph("The subnet — shared with the VMs — is open.")
bad.add_paragraph("Settings were read from Azure using read-only Azure CLI commands.")
bad.add_paragraph("We leverage a robust and seamless approach.")
bad.add_paragraph("Traffic flows Internet → WAF.")
bad.add_paragraph("Public access is enabled; the rule is broad.")
p = bad.add_paragraph()
p.add_run("Impact.").bold = True
p.add_run(" The setting applies to every workspace.")
bad.add_paragraph("This paragraph is repeated word for word to trigger the duplicate check.")
bad.add_paragraph("This paragraph is repeated word for word to trigger the duplicate check.")
bad.add_paragraph("See Section 9.9 for details and Microsoft's guidance.")
bad.save(tmp / "bad.docx")
found = "\n".join(check(tmp / "bad.docx", banned=["contoso-internal"]))
expect("em dash", "em dash" in found)
expect("provenance wording", "provenance" in found and "cli" in found)
expect("AI vocabulary", "leverag" in found and "robust" in found)
expect("arrow in prose", "arrow in prose" in found)
expect("semicolon in prose", "semicolon in prose" in found)
expect("bold lead-in", "bold lead-in" in found)
expect("repeated paragraph", "repeated paragraph" in found)
expect("dangling section reference", "Section 9.9" in found)
expect("straight apostrophe", "straight apostrophe" in found)

# 4. optional: Word finalize + PDF page checks
if os.environ.get("DOCB_WORD") == "1" and os.name == "nt":
    pdf = finalize(clean)
    expect("finalize exported a PDF", pdf and Path(pdf).exists())
    expect("PDF has no empty pages or stranded titles", check(clean, pdf) == [])
    from docx.enum.text import WD_BREAK
    blank = Document()
    blank.add_paragraph("First page text that is long enough to count as content on this page.")
    blank.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    blank.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    blank.add_paragraph("Third page text.")
    blank.save(tmp / "blank.docx")
    bpdf = finalize(tmp / "blank.docx")
    expect("blank page detected", any("page 2 is empty" in i for i in check(tmp / "blank.docx", bpdf)))

sys.exit(0 if ok else 1)
