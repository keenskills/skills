"""Makes the "first draft" for the site's check -> fix showcase from the finished Contoso example.

It puts back five things check exists to catch - an em dash, wording about how the facts were collected,
AI vocabulary, a semicolon in running text and a bold lead-in - and fails unless the real check reports
each one and nothing else, so the draft's findings are genuine docbuilder.py output.

Usage: python3 contoso_draft.py <final.docx> <draft.docx>
Prints JSON: the fixes, what each defect marks in the draft, and both documents as blocks.
Needs DOCBUILDER_DIR = the skill's scripts folder.
"""
import copy
import json
import os
import re
import sys

sys.path.insert(0, os.environ["DOCBUILDER_DIR"])
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402
from docbuilder import check  # noqa: E402

final_path, draft_path = sys.argv[1], sys.argv[2]
doc = Document(final_path)


def para(start):
    return next(p for p in doc.paragraphs if p.text.startswith(start))


def set_text(p, text):
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    p.runs[0].text = text


# Each defect: the paragraph it goes in, the draft text, the part to mark, the check lines it causes, the fix.
DEFECTS = [
    {
        "at": "This document describes",
        "text": "This document describes the proposed network security design for the Contoso pre-production web "
                "applications — the current configuration as of 24-Sep-2026, the target design and the changes "
                "needed to implement it.",
        "mark": "—",
        "lines": r"em dash",
        "fix": "Purpose: the em dash became a plain sentence",
    },
    {
        "at": "This section summarises",
        "text": "This section summarises the configuration of the pre-production subscription, captured with "
                "read-only Azure CLI commands.",
        "mark": "captured with read-only Azure CLI commands",
        "lines": r"provenance",
        "fix": "Current State: how the facts were collected became “as of 24-Sep-2026”",
    },
    {
        "at": "Internet traffic enters",
        "text": "Internet traffic enters through an Application Gateway with WAF. We recommend a robust approach that "
                "leverages private endpoints throughout; the database stays private once the application teams "
                "confirm their dependencies.".replace(";", ";"),
        "mark": "robust approach that leverages",
        "lines": r"AI-sounding",
        "fix": "Proposed Architecture: “a robust approach that leverages” became a plain recommendation",
    },
    {
        "at": None,                      # same paragraph as the one above
        "mark": ";",
        "lines": r"semicolon in prose",
        "fix": "Proposed Architecture: the semicolon split into two sentences",
    },
    {
        "at": "Implementation Plan",
        "insert": ("Note.", " The cloud team applies each change in an agreed change window."),
        "mark": "Note.",
        "lines": r"bold lead-in",
        "fix": "Implementation Plan: the bold lead-in “Note.” was removed",
    },
]

for dft in DEFECTS:
    if dft["at"] is None:
        continue
    p = para(dft["at"])
    if "insert" in dft:                  # a new paragraph under the heading, with a bold first run
        new = copy.deepcopy(para("Internet traffic enters")._p)
        p._p.addnext(new)
        q = Paragraph(new, p._parent)
        set_text(q, dft["insert"][0])
        q.runs[0].bold = True
        q.add_run(dft["insert"][1]).bold = False
    else:
        set_text(p, dft["text"])
doc.save(draft_path)

lines = check(draft_path)
owner = []
for line in lines:
    hits = [i for i, dft in enumerate(DEFECTS) if re.search(dft["lines"], line)]
    if len(hits) != 1:
        sys.exit(f"check line not caused by exactly one defect: {line!r}")
    owner.append(hits[0])
missing = [d["fix"] for i, d in enumerate(DEFECTS) if i not in owner]
if missing:
    sys.exit(f"check did not report: {missing}")


def blocks(path):
    """The body in order, as the site draws it: headings, paragraphs (bold runs kept), bullets, tables."""
    d = Document(path)
    out = []
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, d)
            if not p.text.strip():
                continue
            style = p.style.name
            kind = {"Heading 1": "h1", "Heading 2": "h2", "Heading 3": "h3"}.get(style, "li" if "Bullet" in style else "p")
            runs = [{"text": r.text, **({"bold": True} if r.bold else {})} for r in p.runs if r.text]
            out.append({"kind": kind, "runs": runs})
        elif el.tag == qn("w:tbl"):
            rows = []
            for r in Table(el, d).rows:
                cells, seen = [], set()
                for c in r.cells:              # a merged title row repeats its one cell
                    if id(c._tc) not in seen:
                        seen.add(id(c._tc))
                        cells.append(c.text)
                rows.append(cells)
            title = rows.pop(0)[0] if len(rows[0]) == 1 else None
            out.append({"kind": "table", "title": title, "head": rows[0], "rows": rows[1:]})
    return out


print(json.dumps({
    "fixes": [d["fix"] for d in DEFECTS],
    "marks": [d["mark"] for d in DEFECTS],
    "owner": owner,
    "draft": blocks(draft_path),
    "final": blocks(final_path),
}, ensure_ascii=False))
